"""Builds docs/02_Feature_Requirement_Matrix.xlsx from the register CSV + formula coverage (v0.3 scope column)."""
import csv, json, os, openpyxl
from collections import Counter
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
R=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
H=Font(bold=True,color="FFFFFF"); HF=PatternFill("solid",fgColor="1F3864"); WR=Alignment(wrap_text=True,vertical="top")
def sheet(wb,title,rows,widths):
    ws=wb.create_sheet(title)
    for r in rows: ws.append(r)
    for c in ws[1]: c.font=H; c.fill=HF; c.alignment=WR
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=WR
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions; return ws
wb=openpyxl.Workbook(); wb.remove(wb.active)
req=list(csv.reader(open(f"{R}/engineering/feature_requirement_matrix.csv")))
wsr=sheet(wb,"Requirement_Matrix",req,[10,8,8,48,10,16,13,8,18,22,14,10,16,60,40])
for row in wsr.iter_rows(min_row=2):
    if row[6].value=="OMITTED":
        for c in row: c.font=Font(color="808080",italic=True)
cov=list(csv.reader(open(f"{R}/engineering/formula_coverage_v02.csv")))
ws=sheet(wb,"Formula_Coverage_v02",cov,[9,7,11,40,55,9,16,16,12,40,12])
fills={"COVERED":"C6EFCE","PARTIAL":"FFEB9C","MISSING":"FFC7CE","DIVERGENT":"F4B183"}
for row in ws.iter_rows(min_row=2):
    st=row[8].value
    if st in fills: row[8].fill=PatternFill("solid",fgColor=fills[st])
lg=wb.create_sheet("Legend_Summary"); lg.append(["Field","Meaning"])
for k,v in [("evidence_class","OBS = observed in authorized session; OBS-R = observed and arithmetically reconciled; OBS+PROP = observed entry point, proposed content; DOC = documented; DOC-F = documented formula; HIST = historical; PROP = proposed; NF = not found"),
            ("disposition","IN / CONDITIONAL / NA / OMITTED (OMITTED = marketing function removed by Owner 2026-10-10; kept for traceability, not built)"),("status_before","Coverage of the Sellerboard item in the v0.1 SOW: COVERED / PARTIAL / MISSING / DIVERGENT"),
            ("acceptance","All requirements remain UNACCEPTED until implementation evidence exists"),("v03_scope","Formula coverage rows: IN or OMITTED (marketing, v0.3)"),("version","v0.3 DRAFT — cutoff 2026-10-10; generated from engineering/scope_register.json")]: lg.append([k,v])
lg.append([]); lg.append(["Requirements by disposition"])
for k,v in sorted(Counter(r[6] for r in req[1:]).items()): lg.append([k,v])
lg.append([]); lg.append(["Formula coverage rows by v0.3 scope"])
for k,v in sorted(Counter(r[11] for r in cov[1:]).items()): lg.append([k,v])
lg.append([]); lg.append(["Requirements by release"])
for k,v in sorted(Counter(r[7] for r in req[1:]).items()): lg.append([k,v])
lg.append([]); lg.append(["Sellerboard items by status (before v0.2)"])
for k,v in Counter(r[8] for r in cov[1:] if r[1]!="NEW").items(): lg.append([k,v])
lg.column_dimensions['A'].width=28; lg.column_dimensions['B'].width=120
wb.save(f"{R}/docs/02_Feature_Requirement_Matrix.xlsx"); print("ok",len(req)-1,len(cov)-1)
